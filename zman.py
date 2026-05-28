from datetime import date, timedelta

from zmanim.util.geo_location import GeoLocation
from zmanim.zmanim_calendar import ZmanimCalendar
from pyluach.dates import HebrewDate
import pandas as pd
from pandas import Timestamp

LOCATION = GeoLocation(
    'Savannah, GA',
    32.026,
    -81.108,
    'America/New_York'
)

LATE_CANDLE_LIGHTING_OFFSET = timedelta(minutes=-20)
MINCHA_OFFSET = timedelta(minutes=-15)
TALMUD_CLASS_OFFSET = timedelta(minutes=-90)
SHABBOS_CONCLUDES_OFFSET = timedelta(minutes=45)
SHABBOS_MINCHA_OFFSET = timedelta(minutes=-30)


def get_weekend_dates(start: date, end: date) -> list[date]:
    days = []
    current = start
    while current <= end:
        if current.weekday() in (4, 5, 6):  # Friday, Saturday, Sunday
            days.append(current)
        current += timedelta(days=1)
    return days


def _fmt(dt) -> str:
    if dt is None:
        return "N/A"
    return dt.strftime('%#I:%M %p')


def compute_zmanim(d: date) -> dict:
    cal = ZmanimCalendar(geo_location=LOCATION, date=d)
    weekday = d.weekday()
    data = {
        "_date": d,
        "Date": (
            f"{'Shabbos' if d.weekday() == 5 else d.strftime('%A')}{d.strftime(', %B %d')}\n{HebrewDate.from_pydate(d):%*d %*B}"
        ),
        # "Shkia": _fmt(cal.shkia()),
        # "Plag HaMincha": _fmt(cal.plag_hamincha()),
    }
    if weekday == 4:  # Friday
        data['Early Candle Lighting'] = _fmt(
            Timestamp(cal.plag_hamincha()).ceil('min')
        )
        data['Late Candle Lighting'] = _fmt(
            Timestamp(cal.shkia()).floor('min') + LATE_CANDLE_LIGHTING_OFFSET
        )
        data['Mincha'] = _fmt(
            Timestamp(cal.plag_hamincha()).ceil('min') + MINCHA_OFFSET
        )
    elif weekday == 5:  # Shabbos
        data['Mincha'] = _fmt(
            Timestamp(cal.shkia()).ceil('min') + SHABBOS_MINCHA_OFFSET
        )
        data['Talmud Class'] = _fmt(
            Timestamp(cal.shkia()).ceil('min') + TALMUD_CLASS_OFFSET
        )
        data['Shabbos Concludes'] = _fmt(
            Timestamp(cal.shkia()).ceil('min') + SHABBOS_CONCLUDES_OFFSET
        )
    elif weekday == 6:  # Sunday
        thurs_cal = ZmanimCalendar(
            geo_location=LOCATION, date=d + timedelta(days=4)
        )
        data['Mincha'] = _fmt(
            Timestamp(cal.plag_hamincha()).ceil('min') + MINCHA_OFFSET
        )
        data['Thursday Mincha'] = _fmt(
            Timestamp(thurs_cal.plag_hamincha()).ceil('min') + MINCHA_OFFSET
        )
    return data


def compute_dataframe(rows: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(
        rows,
        columns=[
            '_date',
            'Date',
            'Early Candle Lighting',
            'Talmud Class',
            'Mincha',
            'Late Candle Lighting',
            'Shabbos Concludes',
            'Thursday Mincha',
        ]
    )
    return (
        df.melt(id_vars=['_date', 'Date'], var_name='Zman', value_name='Time')
        .dropna(subset=['Time'])
        .sort_values(by=['_date'], kind='stable')
        .drop(columns=['_date'])
        .reset_index(drop=True)
    )


def generate_excel(rows: list[dict], save_path: str) -> None:
    df = compute_dataframe(rows).set_index(['Date', 'Zman'])

    df.to_excel(save_path, index=True, engine='openpyxl')

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    wb = load_workbook(save_path)
    ws = wb.active
    for col in ws.columns:
        max_len = max((len(str(cell.value)) if cell.value is not None else 0) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 2
    for cell in ws['A']:
        cell.alignment = Alignment(vertical='top', wrap_text=True)
    wb.save(save_path)